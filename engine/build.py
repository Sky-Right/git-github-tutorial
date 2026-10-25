#!/usr/bin/env python3
"""
Markdown -> container.training-style HTML + editable PPTX.

This is intentionally small and hackable. The HTML output uses Remark.js and
theme/workshop.css, which stays close to the original container.training feel.
The PPTX output recreates the same visual language with editable PowerPoint
objects via python-pptx.
"""

from __future__ import annotations

import argparse
import base64
import html
import mimetypes
import os
import re
from pathlib import Path
from typing import Dict, List, Tuple, Optional

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError as exc:
    raise SystemExit(
        "Missing dependency: python-pptx\n"
        "Install it with: pip install -r requirements.txt"
    ) from exc


SLIDE_W = 13.333333
SLIDE_H = 7.5

BLACK = RGBColor(17, 17, 17)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(238, 238, 238)
CODE_GRAY = RGBColor(204, 204, 204)
RED = RGBColor(250, 0, 0)
BLUE = RGBColor(0, 0, 238)


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def split_slides(markdown: str) -> List[str]:
    slides: List[str] = []
    buf: List[str] = []
    for line in markdown.splitlines():
        if line.strip() == "---":
            slides.append("\n".join(buf).strip("\n"))
            buf = []
        else:
            buf.append(line)
    slides.append("\n".join(buf).strip("\n"))
    return [s for s in slides if s.strip()]


def parse_meta(raw_slide: str) -> Tuple[Dict[str, str], str]:
    meta: Dict[str, str] = {}
    lines = raw_slide.splitlines()
    body_start = 0
    for i, line in enumerate(lines):
        if not line.strip():
            body_start = i + 1
            break
        m = re.match(r"^([A-Za-z0-9_-]+)\s*:\s*(.+?)\s*$", line)
        if m:
            meta[m.group(1).lower()] = m.group(2).strip()
            body_start = i + 1
        else:
            break
    return meta, "\n".join(lines[body_start:]).strip("\n")


def first_heading(markdown: str) -> Tuple[Optional[str], str]:
    lines = markdown.splitlines()
    for i, line in enumerate(lines):
        m = re.match(r"^#\s+(.+?)\s*$", line)
        if m:
            rest = "\n".join(lines[:i] + lines[i+1:]).strip("\n")
            return clean_inline(m.group(1)), rest
    return None, markdown


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("**", "").replace("__", "").replace("*", "").replace("_", "")
    text = re.sub(r"\.([A-Za-z0-9_-]+)\[([^\]]*)\]", r"\2", text)
    return text.strip()


def est_lines(text: str, size: float, width_in: float) -> int:
    """Rough wrapped-line estimate for a column text box (CJK chars are
    ~1em wide, Latin roughly half)."""
    cap = max(1.0, width_in * 72.0)
    lines = 1
    used = 0.0
    for ch in text:
        cw = size if ord(ch) > 0x2E80 else size * 0.52
        if used + cw > cap:
            lines += 1
            used = cw
        else:
            used += cw
    return lines


def iter_blocks(markdown: str):
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            yield ("code", "\n".join(code_lines), lang)
            continue

        img = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if img:
            yield ("image", img.group(2), img.group(1))
            i += 1
            continue

        if stripped.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            yield ("quote", " ".join(quote_lines), None)
            continue

        if re.match(r"^#{2,6}\s+", stripped):
            level = len(stripped) - len(stripped.lstrip("#"))
            yield ("subheading", clean_inline(stripped[level:].strip()), level)
            i += 1
            continue

        if re.match(r"^\s*[-*+]\s+", line) or re.match(r"^\s*\d+\.\s+", line):
            items = []
            while i < len(lines):
                l = lines[i]
                if re.match(r"^\s*[-*+]\s+", l):
                    items.append(clean_inline(re.sub(r"^\s*[-*+]\s+", "", l)))
                    i += 1
                elif re.match(r"^\s*\d+\.\s+", l):
                    items.append(clean_inline(re.sub(r"^\s*\d+\.\s+", "", l)))
                    i += 1
                else:
                    break
            yield ("bullets", items, None)
            continue

        # Remark span container: .left[ ... ] spanning multiple lines
        span_open = re.match(r"^\.([A-Za-z0-9_-]+)\[(.*)$", stripped)
        if span_open and not stripped.endswith("]"):
            cls = span_open.group(1)
            inner = [span_open.group(2)] if span_open.group(2).strip() else []
            i += 1
            while i < len(lines):
                if lines[i].strip() == "]":
                    i += 1
                    break
                inner.append(lines[i])
                i += 1
            yield ("span", cls, "\n".join(inner).strip())
            continue

        # Remark span syntax: .class[text]
        remark_span = re.match(r"^\.[A-Za-z0-9_-]+\[(.*)\]$", stripped)
        if remark_span:
            yield ("paragraph", clean_inline(remark_span.group(1)), None)
            i += 1
            continue

        para = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith("```"):
            next_line = lines[i].strip()
            if next_line == "---" or re.match(r"^#{1,6}\s+", next_line):
                break
            if re.match(r"^\s*[-*+]\s+", lines[i]) or re.match(r"^\s*\d+\.\s+", lines[i]):
                break
            if re.match(r"!\[([^\]]*)\]\(([^)]+)\)", next_line):
                break
            para.append(next_line)
            i += 1
        yield ("paragraph", clean_inline(" ".join(para)), None)


def make_html(markdown: str, title: str, template: str) -> str:
    escaped_md = html.escape(markdown, quote=False)
    return template.replace("@@TITLE@@", html.escape(title)).replace("@@MARKDOWN@@", escaped_md)


def make_standalone_html(markdown: str, title: str, template: str,
                         md_path: Path, out_dir: Path) -> str:
    """Bundle the template's local CSS/JS and Markdown images into one HTML file."""
    def local_path(value: str, base: Path) -> Path:
        if re.match(r"[a-zA-Z][a-zA-Z0-9+.-]*:", value) or value.startswith("//"):
            raise ValueError(f"Standalone assets must be local: {value}")
        return (base / value).resolve(strict=True)

    def data_uri(value: str, base: Path) -> str:
        if value.startswith("data:"):
            return value
        path = local_path(value, base)
        mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        return f"data:{mime};base64,{encoded}"

    def stylesheet(match):
        path = local_path(match.group(1), out_dir)
        css = read_text(path)
        if re.search(r"@import\b", css):
            raise ValueError(f"Inline CSS imports before bundling: {path}")
        css = re.sub(r"url\(['\"]?([^)'\"]+)['\"]?\)",
                     lambda m: f'url("{data_uri(m.group(1), path.parent)}")', css)
        return f"<style>{css}</style>"

    def script(match):
        js = read_text(local_path(match.group(1), out_dir))
        js = re.sub(r"</script", r"<\\/script", js, flags=re.I)
        return f"<script>{js}</script>"

    markdown = re.sub(r"(!\[[^\]]*\]\()([^)]+)(\))",
                      lambda m: m.group(1) + data_uri(m.group(2), md_path.parent) + m.group(3), markdown)
    template = re.sub(r'<link\s+rel="stylesheet"\s+href="([^"]+)"\s*/?>', stylesheet, template)
    template = re.sub(r'<script\s+src="([^"]+)"\s*></script>', script, template)
    return make_html(markdown, title, template)


def add_slide_number(slide, index: int, total: int, inverse: bool = False):
    box = slide.shapes.add_textbox(Inches(12.15), Inches(0.05), Inches(1.0), Inches(0.25))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = f"{index}/{total}"
    run.font.name = "Inconsolata"
    run.font.size = Pt(9)
    run.font.color.rgb = WHITE if inverse else BLACK


def add_textbox(slide, text: str, x, y, w, h, font_size=22, bold=False, color=BLACK,
                font_name="Georgia", align=PP_ALIGN.LEFT, fill=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = fill
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_multiline(slide, lines: List[str], x, y, w, h, font_size=22, color=BLACK,
                  font_name="Georgia", bullet=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(5)
        if bullet:
            p.text = f"• {line}"
        else:
            p.text = line
        for run in p.runs:
            run.font.name = font_name
            run.font.size = Pt(font_size)
            run.font.color.rgb = color
    return shape


def add_code_block(slide, code: str, x, y, w, h, font_size=17):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    rect.fill.solid()
    rect.fill.fore_color.rgb = CODE_GRAY
    rect.line.color.rgb = CODE_GRAY
    tf = rect.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.12)
    tf.margin_bottom = Inches(0.12)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = code
    run.font.name = "Consolas"
    run.font.size = Pt(font_size)
    run.font.color.rgb = BLACK
    return rect


def resolve_image(path_text: str, md_path: Path) -> Optional[Path]:
    raw = path_text.strip()
    if raw.startswith("http://") or raw.startswith("https://"):
        return None
    candidate = (md_path.parent / raw).resolve()
    if candidate.exists():
        return candidate
    candidate = (Path.cwd() / raw).resolve()
    if candidate.exists():
        return candidate
    candidate = (md_path.parent.parent / raw).resolve()
    if candidate.exists():
        return candidate
    return None


def add_image_or_placeholder(slide, path_text: str, md_path: Path, x, y, w, h):
    img = resolve_image(path_text, md_path)
    if img is not None:
        try:
            # Fit within the box, preserving aspect ratio (no distortion),
            # centered both ways. SVG is not supported by python-pptx and
            # falls through to the placeholder below.
            from PIL import Image as PILImage
            with PILImage.open(str(img)) as im:
                iw, ih = im.size
            scale = min(w / iw, h / ih)
            fw, fh = iw * scale, ih * scale
            slide.shapes.add_picture(str(img), Inches(x + (w - fw) / 2), Inches(y + (h - fh) / 2),
                                     width=Inches(fw), height=Inches(fh))
            return
        except Exception:
            pass
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = GRAY
    shape.line.color.rgb = BLACK
    tf = shape.text_frame
    tf.text = f"Image: {path_text}"
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.name = "Georgia"
            r.font.size = Pt(18)
            r.font.color.rgb = BLACK


def slide_background(slide, color: RGBColor):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_lab_frame(slide):
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.35), Inches(0.35), Inches(12.65), Inches(6.85))
    rect.fill.solid()
    rect.fill.fore_color.rgb = GRAY
    rect.line.color.rgb = BLACK
    rect.line.width = Pt(1.5)
    add_textbox(slide, ">", 0.55, 0.42, 0.35, 0.3, font_size=20, bold=True, font_name="Consolas")


SPAN_COLUMNS = {
    "left": (1.0, 5.45),
    "right": (6.85, 5.48),
}


def render_span_blocks(slide, span_md: str, x, y, w, md_path: Path, inverse: bool = False):
    """Render the blocks inside a .left[]/.right[] span container as a column.
    Supports subheadings, paragraphs (with **bold** runs), bullets and images
    (fit into the column width, aspect preserved). Returns the y after the
    column so the caller can continue stacking content below it."""
    color = WHITE if inverse else BLACK
    lines = span_md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        img = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if img:
            add_image_or_placeholder(slide, img.group(2), md_path, x, y, w, 3.4)
            y += 3.55
            i += 1
            continue

        if re.match(r"^#{2,6}\s+", stripped):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip().replace("**", "")
            box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.45))
            tf = box.text_frame
            tf.clear()
            tf.word_wrap = True
            run = tf.paragraphs[0].add_run()
            run.text = text
            run.font.name = "Georgia"
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = color
            y += 0.5
            i += 1
            continue

        if re.match(r"^\s*[-*+]\s+", line) or re.match(r"^\s*\d+\.\s+", line):
            items = []
            while i < len(lines) and (re.match(r"^\s*[-*+]\s+", lines[i]) or re.match(r"^\s*\d+\.\s+", lines[i])):
                items.append(re.sub(r"^\s*(?:[-*+]|\d+\.)\s+", "", lines[i]).strip().replace("**", ""))
                i += 1
            for item in items:
                n = est_lines("• " + item, 18, w - 0.25)
                box = slide.shapes.add_textbox(Inches(x + 0.15), Inches(y), Inches(w - 0.15), Inches(0.32 * n + 0.15))
                tf = box.text_frame
                tf.clear()
                tf.word_wrap = True
                run = tf.paragraphs[0].add_run()
                run.text = "• " + item
                run.font.name = "Georgia"
                run.font.size = Pt(18)
                run.font.color.rgb = color
                y += 0.32 * n + 0.1
            continue

        # paragraph (joins consecutive non-blank lines, honors **bold**)
        para = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith("```"):
            nxt = lines[i].strip()
            if (nxt == "]" or re.match(r"^#{1,6}\s+", nxt)
                    or re.match(r"^\s*[-*+]\s+", lines[i]) or re.match(r"^\s*\d+\.\s+", lines[i])
                    or re.match(r"!\[", nxt)):
                break
            para.append(nxt)
            i += 1
        text = " ".join(para)
        n = est_lines(text, 18, w - 0.25)
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.32 * n + 0.15))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        p = tf.paragraphs[0]
        for part in re.split(r"(\*\*[^*]+\*\*)", text):
            if not part:
                continue
            run = p.add_run()
            run.text = part.replace("**", "")
            run.font.name = "Georgia"
            run.font.size = Pt(18)
            run.font.bold = part.startswith("**")
            run.font.color.rgb = color
        y += 0.32 * n + 0.18
    return y


def create_pptx(markdown: str, md_path: Path, out_pptx: Path, title: str):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    blank = prs.slide_layouts[6]
    raw_slides = split_slides(markdown)
    total = len(raw_slides)

    for idx, raw in enumerate(raw_slides, start=1):
        meta, body = parse_meta(raw)
        classes = set(meta.get("class", "").split())
        slide = prs.slides.add_slide(blank)
        inverse = "inverse" in classes
        if inverse:
            slide_background(slide, BLACK)
        else:
            slide_background(slide, WHITE)

        if "lab" in classes:
            add_lab_frame(slide)

        h1, rest = first_heading(body)

        if "pic" in classes:
            blocks = list(iter_blocks(body))
            img_block = next((b for b in blocks if b[0] == "image"), None)
            if h1:
                add_textbox(slide, h1, 0.9, 0.35, 11.55, 0.7, font_size=34, bold=True,
                            color=WHITE if inverse else BLACK, align=PP_ALIGN.CENTER)
                image_y, image_h = 1.2, 5.8
            else:
                image_y, image_h = 0.0, 7.5
            if img_block:
                add_image_or_placeholder(slide, img_block[1], md_path, 0.0, image_y, 13.333, image_h)
            else:
                add_textbox(slide, "PIC slide: add ![alt](path/to/image)", 1.0, 3.2, 11.3, 0.5,
                            font_size=24, color=WHITE if inverse else BLACK, align=PP_ALIGN.CENTER)
            add_slide_number(slide, idx, total, inverse)
            continue

        if "title" in classes:
            y = 2.15
            if h1:
                add_textbox(slide, h1, 0.9, y, 11.55, 0.9, font_size=42, bold=True,
                            color=WHITE if inverse else BLACK, align=PP_ALIGN.CENTER)
                y += 1.0
            text_lines = []
            for btype, val, extra in iter_blocks(rest):
                if btype in ("paragraph", "quote"):
                    if str(val).strip():
                        text_lines.append(str(val).strip())
                elif btype == "bullets":
                    text_lines.extend(val)
                elif btype == "span":
                    for b2, v2, e2 in iter_blocks(str(extra)):
                        if b2 == "paragraph" and str(v2).strip():
                            text_lines.append(str(v2).strip())
            if text_lines:
                add_multiline(slide, text_lines, 1.3, y, 10.75, 1.25, font_size=22,
                              color=WHITE if inverse else BLACK, font_name="Georgia", bullet=False)
                for p in slide.shapes[-1].text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
            add_slide_number(slide, idx, total, inverse)
            continue

        if "quote" in classes:
            if h1:
                add_textbox(slide, h1, 0.9, 0.55, 11.55, 0.7, font_size=34, bold=True,
                            color=WHITE if inverse else BLACK, align=PP_ALIGN.CENTER)
            quote = ""
            for btype, val, extra in iter_blocks(rest):
                if btype == "quote":
                    quote = str(val)
                    break
            if quote:
                add_textbox(slide, f"“{quote}”", 1.2, 2.55, 10.9, 1.3, font_size=31, bold=False,
                            color=WHITE if inverse else BLACK, align=PP_ALIGN.CENTER)
            add_slide_number(slide, idx, total, inverse)
            continue

        # Normal slide
        content_x = 1.0
        content_y = 0.75
        content_w = 11.35
        if h1:
            h1_lines = est_lines(h1, 34, content_w)
            h1_h = max(0.7, h1_lines * 0.64)
            add_textbox(slide, h1, content_x, content_y, content_w, h1_h,
                        font_size=34, bold=True, color=WHITE if inverse else BLACK)
            y = content_y + h1_h + 0.12
        else:
            y = 0.85

        span_base_y = None
        for btype, val, extra in iter_blocks(rest):
            if y > 6.85:
                break
            if btype == "subheading":
                add_textbox(slide, str(val), content_x, y, content_w, 0.42,
                            font_size=24, bold=True, color=WHITE if inverse else BLACK)
                y += 0.5
            elif btype == "paragraph":
                if not str(val).strip():
                    continue
                n = est_lines(str(val), 21, content_w)
                add_textbox(slide, str(val), content_x, y, content_w, 0.32 * n + 0.15,
                            font_size=21, color=WHITE if inverse else BLACK)
                y += 0.34 * n + 0.18
            elif btype == "quote":
                add_textbox(slide, f"“{val}”", content_x + 0.35, y, content_w - 0.7, 0.75,
                            font_size=24, color=WHITE if inverse else BLACK, align=PP_ALIGN.CENTER)
                y += 0.9
            elif btype == "bullets":
                items = list(val)
                height = min(0.4 + len(items) * 0.42, 4.6)
                add_multiline(slide, items, content_x + 0.1, y, content_w - 0.2, height,
                              font_size=22, color=WHITE if inverse else BLACK, bullet=True)
                y += height + 0.2
            elif btype == "code":
                code = str(val)
                line_count = max(1, code.count("\n") + 1)
                height = min(0.55 + line_count * 0.26, 2.65)
                add_code_block(slide, code, content_x, y, content_w, height, font_size=16 if line_count > 5 else 17)
                y += height + 0.25
            elif btype == "span":
                span_cls = str(val)
                if span_cls in SPAN_COLUMNS:
                    # All columns on a slide start at the same top y; the
                    # cursor only advances so content after the columns
                    # lands below the tallest one.
                    if span_base_y is None:
                        span_base_y = y
                    col_x, col_w = SPAN_COLUMNS[span_cls]
                    col_bottom = render_span_blocks(slide, str(extra), col_x, span_base_y, col_w, md_path, inverse=inverse)
                    y = max(y, col_bottom + 0.15)
            elif btype == "image":
                add_image_or_placeholder(slide, str(val), md_path, content_x + 1.2, y, 8.9, 2.35)
                y += 2.55

        add_slide_number(slide, idx, total, inverse)

    # Remove default first empty slide if it exists and is blank? python-pptx starts without slides.
    out_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_pptx)


def main():
    parser = argparse.ArgumentParser(description="Build container.training-style HTML + PPTX from Markdown.")
    parser.add_argument("markdown", type=Path, help="Path to content/slides.md")
    parser.add_argument("--out", type=Path, default=Path("dist"), help="Output directory")
    parser.add_argument("--name", default="deck", help="Output basename, e.g. deck -> deck.html/deck.pptx")
    parser.add_argument("--title", default=None, help="HTML/PPTX title")
    parser.add_argument("--template", type=Path, default=Path("theme/remark.html"), help="Remark HTML template")
    parser.add_argument("--html-only", action="store_true", help="Skip the legacy PPTX renderer")
    parser.add_argument("--standalone", action="store_true", help="Embed local CSS, JS, and Markdown images in HTML")
    args = parser.parse_args()

    markdown_path = args.markdown
    markdown = read_text(markdown_path)
    title = args.title or "Markdown Slides"

    template = read_text(args.template)
    html_out = args.out / f"{args.name}.html"
    pptx_out = args.out / f"{args.name}.pptx"

    rendered = (make_standalone_html(markdown, title, template, markdown_path.resolve(), args.out.resolve())
                if args.standalone else make_html(markdown, title, template))
    write_text(html_out, rendered)
    if not args.html_only:
        create_pptx(markdown, markdown_path.resolve(), pptx_out, title)

    print(f"Created {html_out}")
    if not args.html_only:
        print(f"Created {pptx_out}")


if __name__ == "__main__":
    main()
